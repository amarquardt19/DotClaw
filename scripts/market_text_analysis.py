"""Market-level concept analysis over scraped apartments.com + property-site text.

Outputs an .xlsx with three tabs:
  - Master DF: document frequency across the filtered market
  - Per Neighborhood: concept frequency split by apartments.com neighborhood
  - TF-IDF: top distinctive concepts per property

Usage:
    python3 scripts/market_text_analysis.py --market slc --jsonl data/property_site_text.jsonl --out data/slc_concepts.xlsx
"""

import argparse
import json
import math
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

PRECLUDED_PAGES = re.compile(r"floorplan|floor-plan|floor_plans|availability|pricing", re.I)

STOPWORDS = {
    "the", "a", "an", "and", "or", "of", "to", "in", "for", "on", "at", "with",
    "your", "our", "you", "we", "is", "are", "new", "home", "start", "view",
    "floor", "plan", "plans", "bed", "bath", "rent", "apartments", "apartment",
    "community", "welcome", "contact", "us", "apply", "schedule", "tour", "gallery",
    "amenities", "available", "availability", "call", "coming", "soon", "sq", "ft",
    "month", "de", "see", "find", "show", "more", "read", "less", "close", "clear",
    "now", "today", "back", "list", "details", "hours", "street", "west", "east",
    "north", "south", "living", "style", "modern", "boise", "idaho", "salt lake",
    "slc", "utah", "ut",
}

JUNK_PATTERNS = [
    r"Skip to main content", r"[Cc]ookie\S*", r"Map [Dd]ata \S*[©\w]* ?\S*",
    r"Keyboard shortcuts", r"Terms", r"[Zz]oom", r"[Vv]iew on [Mm]ap", r"Pet Policy",
    r"Office [Hh]ours", r"I Agree", r"I Accept All\S*", r"I Reject All\S*",
    r"[Cc]ookie [Pp]references\S*", r"Floorplan layout\S*", r"Base Rent",
    r"Move [Ii]n [Dd]ate", r"Filters?\b", r"Walkscore for \w+", r"Walk [Ss]core",
    r"Bike [Ss]core", r"Walkscore", r"OpenStreetMap\S*", r"OpenFreeMap\S*", r"Google",
    r"[←→↑↓]", r"Jump \w+[^ ]*", r"[Ll]oad [Aa]ll", r"Privacy Policy\S*",
    r"How We Use Your Data\S*", r"All [Rr]ights [Rr]eserved\S*", r"[Ww]eb [Dd]esign\S*",
    r"Contact Method Popup\S*", r"Accessibility Screen-?Reader\S* Guide\S*",
    r"Click to toggle[^.]*", r"Report a map error\S*", r"AudioEye\S*",
    r"SuMoTuWeThFrSa\S*", r"BESbswy\S*", r"KeyAction\S*", r"Reset [Ff]ilters\S*",
    r"Toggle between metric[^.]*", r"[A-Z][a-z]+jump \w+",
]

NOISE_CONCEPTS = {
    "this", "that", "these", "they", "there", "when", "whether", "what", "here",
    "just", "every", "any", "only", "once", "some", "one", "please", "thank",
    "explore", "enjoy", "experience", "discover", "imagine", "take", "move", "click",
    "book", "browse", "visit", "choose", "check", "share", "spend", "walk", "live",
    "come", "go", "open", "use", "feel", "see", "find", "enter", "now", "today",
    "day", "days", "applicants", "oops", "escape", "nestled", "surrounded",
    "unparalleled", "beautiful", "spacious", "luxury", "convenient", "located",
    "designed", "neighborhood", "center", "park", "pool", "private", "reserved",
    "location", "great", "love", "half", "there's", "yours", "join", "kickstart",
    "services", "staff", "updated", "brand", "good", "area", "kitchen", "sinks",
    "cable", "magazine", "ranked", "named", "combining", "old", "pair", "s paradise",
    "fulfilled", "story", "minute", "minutes", "bathrooms any", "bedrooms any",
    "bed", "bath", "done", "clear", "start", "local", "view units", "square feet",
    "sq ft", "dryer", "washer", "floor plans", "bedroom", "bathrooms", "de", "square",
    "road", "blvd", "boulevard", "street", "st", "ave", "rd", "city", "state", "country",
    "from", "with", "for", "each", "more", "sign in", "tap", "swipe", "schedule",
    "apply", "pricing", "availability", "amenities", "features", "portfolio: skip",
    "unavailable |", "new window", "privacy preferences", "cookies", "feedback",
}

NOISE_CONCEPT_SUBSTRINGS = [
    "jump", "zoom", "cookie", "policy", "preferences", "rights reserved",
    "web design", "openmaptiles", "hours", "pm ", "am-", "asked questions",
    "popup", "shortcut", "keyboard", "scroll", "audioeye", "toggle", "drag",
    "page up", "page down", "mon-fri", "previous", "move left", "move right",
    "home ", "day", "week", "month", "sunday", "monday", "friday", "saturday",
]

CONCEPT_RE = re.compile(
    r"\b((?:[A-Z0-9&'][A-Za-z0-9&'\-]*[a-z]|[A-Z][A-Za-z0-9&'\-]*)(?:\s+(?:"
    r"[A-Z][A-Za-z0-9&'\-]*|at|of|the|de|and|on"
    r")(?!,|\.))*\b)",
)

ATTRIBUTE_PATTERNS = {
    "generic": [
        (r"near downtown|downtown .*|close to downtown|walking distance to downtown|minutes? from downtown|downtown core|heart of (?:the )?(?:city|downtown)", "downtown proximity"),
        (r"close to the (?:freeway|highway)|easy (?:freeway|highway) access|convenient (?:freeway|highway)|easy access to (?:the )?(?:i-?\d+|freeway|highway)|minutes? from the (?:freeway|highway)", "highway access"),
        (r"in[- ]unit (?:laundry|washer)|washer (?:and|&) dryer in (?:each|every)|washer/dryer (?:included|in (?:each|every) (?:home|unit|apartment))", "in-unit laundry"),
        (r"pet[- ]friendly|dog park|pet (?:spa|detailing|centric)", "pet-friendly"),
        (r"ev (?:charging|charger)|electric vehicle charg", "EV charging"),
        (r"smart home|dwelo", "smart home"),
        (r"(?:rooftop|roof ?deck|skyline|panoramic|city|mountain|foothill|river)s? view", "views"),
        (r"coworking|co[- ]working|remote work|office space", "coworking"),
        (r"24[- ]?hour (?:fitness|gym)", "24h gym"),
        (r"resort[- ]style pool|heated pool|year[- ]round (?:pool|spa)", "resort/heated pool"),
        (r"(?:covered|garage|private) parking|attached garage|reserved (?:spot|parking)", "parking (covered/garage)"),
        (r"walkscore|walk score|walker'?s paradise|walkable", "walkability"),
        (r"newly renovated|recently renovated|renovated (?:in|units)|2019[-–]202\d renovated", "renovated"),
        (r"smoke[- ]free", "smoke-free"),
        (r"controlled access|gated", "controlled access"),
        (r"furnished|corporate housing", "furnished/corporate"),
        (r"affordable|income[- ]restricted|affordable program", "affordable/income-restricted"),
        (r"\briver\b", "river mention"),
        (r"\bfoothills?\b", "foothills"),
    ],
    "slc": [
        (r"university of utah|\bu of u\b|utes\b", "University of Utah"),
        (r"\bi-?15\b", "I-15"),
        (r"\bi-?80\b", "I-80"),
        (r"\btrax\b|utah transit authority|\buta\b", "TRAX/UTA"),
        (r"front.?runner|commuter rail", "FrontRunner"),
        (r"big cottonwood canyon|little cottonwood canyon|\bski\b|snowbird|alta|brighton|solitude", "ski access"),
        (r"temple square|\btemple\b", "temple"),
        (r"university hospital|intermountain|primary children", "hospital/healthcare"),
        (r"rice-eccles|real salt lake|\brsl\b|stadium", "stadium/RSL"),
        (r"silicon slopes|lehi|draper tech", "Silicon Slopes"),
    ],
    "boise": [
        (r"(?:on|near|access to|adjacent to) (?:the )?(?:boise river )?greenbelt", "greenbelt access"),
        (r"\bmicron\b", "Micron (employer)"),
        (r"\bi-?84\b", "I-84"),
        (r"bogu?s basin", "Bogus Basin"),
        (r"boise state|\bbsu\b", "Boise State"),
    ],
}


def clean_page_text(text):
    for pat in JUNK_PATTERNS:
        text = re.sub(pat, " ", text, flags=re.I)
    return text


def attribute_res_for(market):
    patterns = list(ATTRIBUTE_PATTERNS["generic"])
    if market in ATTRIBUTE_PATTERNS:
        patterns.extend(ATTRIBUTE_PATTERNS[market])
    return [(re.compile(p, re.I), label) for p, label in patterns]


def extract_concepts(text, attribute_res):
    text = clean_page_text(text)
    concepts = []
    for cre, label in attribute_res:
        if cre.search(text):
            concepts.append(label)
    for m in CONCEPT_RE.finditer(text):
        raw = " ".join(m.group(0).split())
        drop_props = re.sub(r"'s\b", "", raw)
        words = drop_props.split()
        if len(words) == 1 and len(words[0]) < 4:
            continue
        lowered = [re.sub(r"['’]s$", "", w.lower()) for w in words]
        while lowered and lowered[0] in STOPWORDS:
            lowered.pop(0)
        while lowered and lowered[-1] in STOPWORDS:
            lowered.pop()
        if not lowered or len(lowered) > 6:
            continue
        if all(w in STOPWORDS for w in lowered):
            continue
        concept = " ".join(lowered)
        if concept in NOISE_CONCEPTS:
            continue
        if any(sub in concept for sub in NOISE_CONCEPT_SUBSTRINGS):
            continue
        concepts.append(concept)
    return concepts


def name_words(deal_name):
    name = deal_name.split(",")[0]
    name = re.sub(r"\d+.*", "", name)
    return {w for w in re.findall(r"[a-z]+", name.lower()) if len(w) > 2}


def load_docs(jsonl_path, market):
    attribute_res = attribute_res_for(market)
    docs = {}
    with open(jsonl_path) as f:
        for line in f:
            if not line.strip():
                continue
            rec = json.loads(line)
            name = rec.get("deal_name") or ""
            city = (rec.get("city") or "").lower()
            neighborhood = (rec.get("neighborhood") or "").lower()
            text_blob = " ".join([
                rec.get("apartments_text") or "",
                " ".join(p.get("text", "") for p in rec.get("pages", [])),
            ])
            if not text_blob.strip():
                continue
            if market.lower() not in name.lower() and market.lower() not in city and market.lower() not in neighborhood:
                # SLC special-case: accept 'salt lake'
                if not (market.lower() == "slc" and "salt lake" in name.lower()):
                    continue
            hood = rec.get("neighborhood") or "Unknown"
            concepts = []
            own = name_words(name)
            for page in rec.get("pages", []):
                page_name = str(page.get("page", "")).lower()
                if PRECLUDED_PAGES.search(page_name):
                    continue
                page_concepts = extract_concepts(page.get("text", ""), attribute_res)
                page_concepts = [c for c in page_concepts if not own.intersection(c.split())]
                concepts.extend(page_concepts)
            if rec.get("apartments_text"):
                apt_concepts = extract_concepts(rec["apartments_text"], attribute_res)
                apt_concepts = [c for c in apt_concepts if not own.intersection(c.split())]
                concepts.extend(apt_concepts)
            docs[name] = {"neighborhood": hood, "concepts": concepts}
    return docs


def build_master_df(docs):
    doc_freq = Counter()
    term_freq = {}
    for name, data in docs.items():
        tf = Counter(data["concepts"])
        term_freq[name] = tf
        for c in set(data["concepts"]):
            doc_freq[c] += 1
    rows = []
    for concept, df_cnt in sorted(doc_freq.items(), key=lambda kv: (-kv[1], kv[0])):
        if df_cnt < 2:
            continue
        total = sum(term_freq[n][concept] for n in docs)
        rows.append({"Concept": concept, "# Properties": df_cnt, "Count": total})
    return rows


def build_neighborhood_df(docs):
    neighborhood_concept_counts = defaultdict(Counter)  # neighborhood -> concept -> total mentions
    neighborhood_property_counts = defaultdict(Counter)  # neighborhood -> concept -> # properties
    for name, data in docs.items():
        hood = data["neighborhood"]
        seen = set(data["concepts"])
        tf = Counter(data["concepts"])
        for c, cnt in tf.items():
            neighborhood_concept_counts[hood][c] += cnt
            neighborhood_property_counts[hood][c] += 1
    rows = []
    for hood, counts in neighborhood_concept_counts.items():
        for concept in sorted(counts, key=lambda c: -counts[c]):
            rows.append({
                "Neighborhood": hood,
                "Concept": concept,
                "# Properties": neighborhood_property_counts[hood][concept],
                "Count": counts[concept],
            })
    return rows


def build_tfidf_rows(docs):
    doc_freq = Counter()
    term_freq = {}
    for name, data in docs.items():
        tf = Counter(data["concepts"])
        term_freq[name] = tf
        for c in set(data["concepts"]):
            doc_freq[c] += 1
    n_docs = len(docs)
    rows = []
    for name in sorted(docs):
        tf = term_freq[name]
        if not tf:
            continue
        scores = {}
        for concept, cnt in tf.items():
            df_cnt = doc_freq[concept]
            idf = math.log((n_docs + 1) / (df_cnt + 1)) + 1
            scores[concept] = (1 + math.log(cnt)) * idf
        top = sorted(scores.items(), key=lambda kv: -kv[1])[:10]
        for rank, (concept, score) in enumerate(top, 1):
            rows.append({
                "Property": name,
                "Rank": rank,
                "Concept": concept,
                "Count": tf[concept],
                "TF-IDF": round(score, 3),
            })
    return rows


def _write_sheet(ws, rows, headers):
    ws.append(headers)
    for col, _ in enumerate(headers, 1):
        ws.cell(row=1, column=col).font = Font(bold=True)
    for r in rows:
        ws.append([r[h] for h in headers])


def write_xlsx(out_path, master_rows, hood_rows, tfidf_rows):
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Master DF"
    _write_sheet(ws1, master_rows, ["Concept", "# Properties", "Count"])

    ws2 = wb.create_sheet("Per Neighborhood")
    _write_sheet(ws2, hood_rows, ["Neighborhood", "Concept", "# Properties", "Count"])

    ws3 = wb.create_sheet("TF-IDF")
    _write_sheet(ws3, tfidf_rows, ["Property", "Rank", "Concept", "Count", "TF-IDF"])

    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--market", required=True, help="market slug: slc, boise, etc.")
    parser.add_argument("--jsonl", default="data/property_site_text.jsonl")
    parser.add_argument("--out", default="data/market_concepts.xlsx")
    args = parser.parse_args()

    jsonl_path = Path(args.jsonl)
    if not jsonl_path.exists():
        raise FileNotFoundError(jsonl_path)

    docs = load_docs(jsonl_path, args.market)
    print(f"{len(docs)} properties matched market '{args.market}'")

    master_rows = build_master_df(docs)
    hood_rows = build_neighborhood_df(docs)
    tfidf_rows = build_tfidf_rows(docs)

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    write_xlsx(out_path, master_rows, hood_rows, tfidf_rows)
    print(f"wrote {out_path}")


if __name__ == "__main__":
    main()
