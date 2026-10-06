import argparse
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import urlparse

from openpyxl import load_workbook

DEFAULT_URL = "https://www.apartments.com/apartments/san-mateo-county-ca/?bb=0oqxk92g0Ox_p79kO"
DEFAULT_WORKBOOK = str(Path.home() / "Downloads" / "dealsifta apartments.com.xlsx")
DEFAULT_SHEET = None


def _norm_text(val: Any) -> str:
    if val is None:
        return ""
    return " ".join(str(val).split()).strip()


def _build_header_map(ws) -> Dict[str, int]:
    out = {}
    for col in range(1, ws.max_column + 1):
        key = ws.cell(1, col).value
        if key is None:
            continue
        out[str(key).strip().lower()] = col
    return out


def _open_workbook(path: str, sheet_name: Optional[str]):
    wb = load_workbook(path)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    return wb, ws


def _is_property_url(url: str) -> bool:
    if not url:
        return False
    try:
        parsed = urlparse(url)
    except Exception:
        return False
    if parsed.netloc not in {"www.apartments.com", "apartments.com"}:
        return False
    path = (parsed.path or "").strip("/")
    if not path:
        return False
    if path.lower().startswith("apartments/"):
        return False
    parts = [p for p in path.split("/") if p]
    if len(parts) != 2:
        return False
    slug, listing_id = parts
    if not slug or not listing_id:
        return False
    return True


def _collect_listing_urls(page) -> List[str]:
    seen = set()
    results: List[str] = []
    stagnant = 0
    prev_count = 0
    for _ in range(30):
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
        time.sleep(1.2)
        urls = page.evaluate(
            """() => {
                const out = [];
                document.querySelectorAll('a[href], a.property-link').forEach((a) => {
                    const href = a.href || '';
                    if (!href.includes('apartments.com/')) return;
                    if (href.includes('/apartments/') || href.includes('#')) return;
                    out.push(href);
                });
                return out;
            }"""
        ) or []
        for url in urls:
            if url in seen:
                continue
            if not _is_property_url(url):
                continue
            seen.add(url)
            results.append(url)
        cur_count = len(results)
        if cur_count == prev_count:
            stagnant += 1
        else:
            stagnant = 0
        prev_count = cur_count
        if stagnant >= 4:
            break
    return results


def _capture_selected_page_text(page) -> str:
    try:
        page.bring_to_front()
    except Exception:
        pass
    try:
        page.evaluate("window.scrollTo(0, 0)")
        time.sleep(0.5)
        page.evaluate(
            """() => {
                if (document.activeElement && document.activeElement.blur) {
                    document.activeElement.blur();
                }
                if (document.body && document.body.focus) {
                    if (!document.body.hasAttribute('tabindex')) {
                        document.body.setAttribute('tabindex', '-1');
                    }
                    document.body.focus();
                }
            }"""
        )
        time.sleep(0.3)
        page.keyboard.press("Meta+A")
        time.sleep(0.5)
        selected = page.evaluate("() => window.getSelection ? window.getSelection().toString() : ''") or ""
        if selected.strip():
            return selected
    except Exception:
        pass
    return page.evaluate("() => document.body && document.body.innerText ? document.body.innerText : ''") or ""


def _upsert_notes_raw(ws, header_map: Dict[str, int], source_url: str, notes_raw: str) -> None:
    source_col = header_map.get("source_mailbox")
    notes_raw_col = header_map.get("notes_raw")
    if not source_col or not notes_raw_col:
        raise RuntimeError("Workbook must contain 'source_mailbox' and 'notes_raw' columns")
    target_row = None
    for row_idx in range(2, ws.max_row + 1):
        existing_source = _norm_text(ws.cell(row_idx, source_col).value)
        if existing_source == source_url:
            target_row = row_idx
            break
    if target_row is None:
        target_row = ws.max_row + 1
        ws.cell(target_row, source_col).value = source_url
    ws.cell(target_row, notes_raw_col).value = notes_raw[:30000] if notes_raw else None


def main() -> None:
    parser = argparse.ArgumentParser(description="Load Apartments.com property-page select-all text into notes_raw")
    parser.add_argument("--url", default=DEFAULT_URL)
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--max", type=int, default=20)
    parser.add_argument("--skip", type=int, default=0)
    parser.add_argument("--headless", action="store_true")
    parser.add_argument("--cdp-url", default=None)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        raise SystemExit("Playwright not installed")

    if not Path(args.workbook).exists():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    wb, ws = _open_workbook(args.workbook, args.sheet)
    header_map = _build_header_map(ws)

    with sync_playwright() as p:
        close_browser = False
        if args.cdp_url:
            browser = p.chromium.connect_over_cdp(args.cdp_url)
            context = browser.contexts[0] if browser.contexts else browser.new_context(viewport={"width": 1600, "height": 1200})
            page = context.pages[0] if context.pages else context.new_page()
        else:
            browser = p.chromium.launch(headless=args.headless)
            context = browser.new_context(viewport={"width": 1600, "height": 1200})
            page = context.new_page()
            close_browser = True

        if _is_property_url(args.url):
            listings = [args.url]
        else:
            page.goto(args.url, timeout=30000, wait_until="domcontentloaded")
            time.sleep(5)
            listings = _collect_listing_urls(page)

        sliced = listings[args.skip:args.skip + args.max]
        print(f"Found {len(listings)} listings, processing {len(sliced)}", flush=True)
        written = 0
        skipped = 0
        for idx, url in enumerate(sliced, start=1):
            detail_page = context.new_page()
            try:
                print(f"[{idx}/{len(sliced)}] {url}", flush=True)
                detail_page.goto(url, timeout=30000, wait_until="domcontentloaded")
                time.sleep(4)
                notes_raw = _capture_selected_page_text(detail_page)
                if not notes_raw.strip():
                    print("    skip empty notes_raw", flush=True)
                    skipped += 1
                    continue
                if not args.dry_run:
                    _upsert_notes_raw(ws, header_map, url, notes_raw)
                print(f"    ok notes_raw_len={len(notes_raw)}", flush=True)
                written += 1
            except Exception as exc:
                print(f"    failed: {exc}", flush=True)
                skipped += 1
            finally:
                detail_page.close()

        if close_browser:
            browser.close()

    if not args.dry_run:
        wb.save(args.workbook)
    print(f"Done. written={written} skipped={skipped} dry_run={args.dry_run}", flush=True)


if __name__ == "__main__":
    main()
